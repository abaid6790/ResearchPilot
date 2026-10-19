"""Phase 10: quality scoring and exports (files are opened and inspected, not just created)."""

from __future__ import annotations

import io
import zipfile

import fitz  # PyMuPDF
import pytest
from docx import Document
from lxml import html as lxml_html
from sqlalchemy import select

from agents.schemas import DimensionOpinion, QualityOpinion
from core.audit import AuditEvent
from extensions import db
from reports import exports, quality
from reports.models import QualityEvaluation, ReportDocument, ReportSection
from tests.conftest import login
from tests.fake_llm import FakeLLM
from tests.test_phase9_editor import ac, alice, bob, build, csrf_post, url  # noqa: F401


def finish(g, *, extra_sections=0):
    """Make the fixture graph a complete, checked report: limitations, claim support, validation flags."""
    report, job = g["report"], g["job"]
    for old in db.session.scalars(
        select(ReportSection).where(ReportSection.report_id == report.id, ReportSection.key == "references")
    ):
        db.session.delete(old)  # the fixture starts with an empty one; the pipeline writes the real list
    db.session.flush()
    g["claim"].support_status = "supported"
    job.checkpoint = {**job.checkpoint, "validation_complete": True, "review": {"issues": []}}
    for i in range(extra_sections):
        db.session.add(ReportSection(report_id=report.id, key=f"extra{i}", heading=f"Extra {i}", position=5 + i,
                                     body=f"More on topic {i} [@S{g['sources'][0].id}]."))  # fmt: skip
    db.session.add(ReportSection(report_id=report.id, key="limitations", heading="Limitations and uncertainty",
                                 position=20, body="- Only some sources were read in full."))  # fmt: skip
    db.session.add(
        ReportSection(report_id=report.id, key="references", heading="References", position=21, body="")
    )
    db.session.commit()
    return g


@pytest.fixture()
def g(alice):
    g = build(alice)
    db.session.delete(
        db.session.scalars(select(ReportSection).where(ReportSection.key == "references")).one()
    )
    db.session.commit()
    return finish(g)


def section(g, key="overview"):
    return db.session.scalar(
        select(ReportSection).where(ReportSection.report_id == g["report"].id, ReportSection.key == key)
    )


# ------------------------------------------------------------------ quality scoring
def test_scores_follow_the_documented_formulas(g):
    ev = quality.compute(g["report"])
    got = {s.dimension: s.score for s in ev.scores}
    assert got["citation_integrity"] == 100 and got["question_coverage"] == 100 and got["completeness"] == 100
    assert got["source_diversity"] == 100 and got["limitations"] == 100 and got["evidence_coverage"] == 100
    assert (
        got["source_relevance"] is None
        and got["contradiction_handling"] is None
        and got["clarity_structure"] is None
    )
    assert got["overall"] == 100 and ev.gates_passed
    # a quoted passage that is in no cited source costs 15 points and fails the citation gate
    section(g).body = f'He said "this sentence is nowhere in any source" [@S{g["sources"][0].id}].'
    db.session.commit()
    ev = quality.compute(g["report"])
    assert ev.get("citation_integrity").score == 85 and not ev.gates_passed
    assert any(w["kind"] == "unverified_quote" for w in ev.warnings)


def test_invalid_citation_scores_by_share(g):
    sid = g["sources"][0].id
    section(g).body = f"One [@S{sid}] two [@S999999] three [@S{sid}] four [@S888888]."
    db.session.commit()
    ci = quality.compute(g["report"]).get("citation_integrity")
    assert ci.score == 50 and ci.details["invalid_source_ids"] == [888888, 999999] and ci.passed_gate is False


def test_relevance_and_diversity_use_real_stored_values(g):
    g["sources"][0].relevance_score = 0.8
    db.session.commit()
    ev = quality.compute(g["report"])
    assert ev.get("source_relevance").score == 80
    section(g).body += f" And [@S{g['sources'][1].id}] [@S{g['sources'][2].id}]."
    db.session.commit()
    div = quality.compute(g["report"]).get("source_diversity")
    # 3 cited sources: 3 domains / min(5,3)=1.0 ; 3 categories / min(3,3)=1.0
    assert div.score == 100 and div.details["cited_sources"] == 3


def test_contradictions_not_cited_lower_score(g):
    from research.models import ClaimSource

    db.session.add(
        ClaimSource(
            claim_id=g["claim"].id, source_id=g["sources"][2].id, job_id=g["job"].id, relation="contradicts"
        )
    )
    db.session.commit()
    c = quality.compute(g["report"]).get("contradiction_handling")
    assert c.score == 0
    section(g).body += f" Others disagree [@S{g['sources'][2].id}]."
    db.session.commit()
    assert quality.compute(g["report"]).get("contradiction_handling").score == 100


def test_gates_fail_for_missing_pieces(g):
    assert quality.compute(g["report"]).gates_passed
    g["claim"].support_status = "insufficient"
    db.session.commit()
    assert not next(x for x in quality.compute(g["report"]).gates if x.key == "claims").passed
    g["claim"].support_status = "supported"
    db.session.delete(section(g, "limitations"))
    db.session.commit()
    ev = quality.compute(g["report"])
    assert not next(x for x in ev.gates if x.key == "limitations").passed
    assert any("no Limitations" in w["text"] for w in ev.warnings)


def test_unreviewed_critic_and_open_issues_fail_the_review_gate(g):
    g["job"].checkpoint = {**g["job"].checkpoint, "review": {"unavailable": True}}
    db.session.commit()
    assert not next(x for x in quality.compute(g["report"]).gates if x.key == "review").passed
    g["job"].checkpoint = {
        **g["job"].checkpoint,
        "review": {
            "issues": [
                {
                    "severity": "major",
                    "section_key": "overview",
                    "problem": "p",
                    "kind": "unclear",
                    "suggestion": "",
                }
            ]
        },
    }
    db.session.commit()
    assert not next(x for x in quality.compute(g["report"]).gates if x.key == "review").passed
    section(g).edited_by_user = True  # edited sections are not re-reviewed: warned about, not failed
    db.session.commit()
    ev = quality.compute(g["report"])
    assert next(x for x in ev.gates if x.key == "review").passed
    assert any(w["kind"] == "edited" for w in ev.warnings)


def test_persist_stores_rows_and_updates_checked_flag(g):
    report = g["report"]
    run_id, ev = quality.persist(report)
    db.session.commit()
    rows = db.session.scalars(select(QualityEvaluation).where(QualityEvaluation.run_id == run_id)).all()
    assert {r.dimension for r in rows} == {s.dimension for s in ev.scores}
    assert all(r.evaluator == "deterministic" and r.method for r in rows)
    na = [r for r in rows if r.status == "not_assessed"]
    assert na and all(r.score is None for r in na)
    overall = next(r for r in rows if r.dimension == "overall")
    assert overall.passed_gate is True and overall.details["gates"]
    assert db.session.get(ReportDocument, report.id).export_ready is True and report.last_validated_at


def test_llm_assessment_is_stored_separately_and_never_moves_the_gates(app, ac, g, alice):
    app.extensions["llm_router"] = FakeLLM(critic=lambda r, s: QualityOpinion(opinions=[
        DimensionOpinion(dimension="clarity_structure", score=72, rationale="Clear enough."),
        DimensionOpinion(dimension="completeness", score=40, rationale="Thin.")]))  # fmt: skip
    before = quality.compute(g["report"]).gates_passed
    rows = quality.assess_with_llm(g["report"], alice)
    db.session.commit()
    assert {r.evaluator for r in rows} == {"llm"} and {r.dimension for r in rows} == {
        "clarity_structure",
        "completeness",
    }
    assert quality.compute(g["report"]).gates_passed == before
    assert not quality.latest_rows(g["report"], "deterministic")
    html = ac.get(url(g, "/report/quality")).get_data(as_text=True)
    assert "AI assessment (an opinion)" in html and "Clear enough." in html and "Thin." in html


def test_quality_page_and_run_endpoint(ac, g):
    html = ac.get(url(g, "/report/quality")).get_data(as_text=True)
    assert "Calculated score" in html and "Quality gates" in html and "How it is calculated" in html
    assert "not assessed" in html and "<script>" not in html
    r = csrf_post(ac, g, url(g, "/report/quality/run"))
    assert r.status_code == 302
    assert db.session.get(ReportDocument, g["report"].id).export_ready is True
    assert ac.post(url(g, "/report/quality/run")).status_code == 400  # CSRF


def test_edit_then_rerun_checks_restores_checked_flag(ac, g):
    csrf_post(ac, g, url(g, "/report/quality/run"))
    sid = g["sources"][0].id
    csrf_post(
        ac, g, url(g, "/report/edit/overview/save"), {"body": f"Edited and cited [@S{sid}].", "version": 1}
    )
    assert db.session.get(ReportDocument, g["report"].id).export_ready is False
    csrf_post(ac, g, url(g, "/report/quality/run"))
    db.session.expire_all()
    assert db.session.get(ReportDocument, g["report"].id).export_ready is True
    # an edit that breaks a citation is caught on re-run
    csrf_post(ac, g, url(g, "/report/edit/overview/save"), {"body": "Broken [@S999999].", "version": 2})
    csrf_post(ac, g, url(g, "/report/quality/run"))
    db.session.expire_all()
    assert db.session.get(ReportDocument, g["report"].id).export_ready is False


# ------------------------------------------------------------------ exports
def get(ac, g, fmt):
    return ac.get(url(g, f"/report/export/{fmt}"))


def pdf_text(data: bytes) -> tuple[str, fitz.Document]:
    pdf = fitz.open(stream=data, filetype="pdf")
    return "\n".join(page.get_text() for page in pdf), pdf


def test_markdown_export(ac, g):
    r = get(ac, g, "md")
    assert r.status_code == 200 and r.mimetype == "text/markdown"
    assert 'filename="Report-a-' in r.headers["Content-Disposition"] and r.headers[
        "Content-Disposition"
    ].endswith('.md"')
    text = r.get_data(as_text=True)
    assert text.startswith("# Report a") and "## Research scope" in text and "## Methodology" in text
    assert "COP of 3 to 4 in mild weather [1]." in text and "[@S" not in text
    assert "1. Heat pump guide a - <https://energy.gov/a0>" in text and "## Quality assessment" in text
    assert "NOT FULLY VERIFIED" not in text  # gates pass on the finished fixture


def test_unverified_reports_say_so_in_every_format(ac, g):
    section(g).body = "Broken [@S999999]."
    db.session.commit()
    assert "NOT FULLY VERIFIED" in get(ac, g, "md").get_data(as_text=True)
    assert "NOT FULLY VERIFIED" in get(ac, g, "html").get_data(as_text=True)
    assert "NOT FULLY VERIFIED" in "\n".join(
        p.text for p in Document(io.BytesIO(get(ac, g, "docx").data)).paragraphs
    )
    assert "NOT FULLY VERIFIED" in pdf_text(get(ac, g, "pdf").data)[0].replace("\n", " ")


def test_html_export_is_escaped_and_linked(ac, g):
    section(
        g
    ).body = f"<script>alert(1)</script> <img src=x onerror=alert(2)> **bold** [@S{g['sources'][0].id}] [x](javascript:alert(3))"
    g["sources"][1].url = "javascript:alert(4)"
    db.session.commit()
    r = get(ac, g, "html")
    assert r.mimetype == "text/html" and "sandbox" in r.headers["Content-Security-Policy"]
    assert r.headers["Content-Disposition"].startswith("attachment")
    raw = r.get_data(as_text=True)
    assert "<script>alert" not in raw and "onerror=" not in raw.replace(
        "&lt;img src=x onerror=alert(2)&gt;", ""
    )
    tree = lxml_html.fromstring(raw)
    assert not tree.xpath("//script") and not tree.xpath("//img") and not tree.xpath("//*[@onerror]")
    assert not [h for h in tree.xpath("//a/@href") if h.lower().startswith("javascript")]
    assert tree.xpath('//a[@href="#ref-1"]') and tree.xpath('//li[@id="ref-1"]')
    ext = tree.xpath('//li[@id="ref-1"]//a')[0]
    assert ext.get("href") == "https://energy.gov/a0" and "noopener" in ext.get("rel")
    assert tree.xpath("//h1")[0].text == "Report a"


def test_docx_export_is_a_real_document(ac, g):
    r = get(ac, g, "docx")
    assert r.mimetype == exports.FORMATS["docx"][0].split(";")[0]
    assert zipfile.ZipFile(io.BytesIO(r.data)).testzip() is None
    d = Document(io.BytesIO(r.data))
    headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert headings[:3] == ["Research scope", "Methodology", "Overview"] and "References" in headings
    assert d.paragraphs[5].text == "Report a" or any(
        p.text == "Report a" and p.style.name == "Title" for p in d.paragraphs
    )
    xml = d.element.xml
    assert 'w:anchor="ref1"' in xml and 'w:name="ref1"' in xml  # citation -> reference bookmark
    assert "https://energy.gov/a0" in {rel.target_ref for rel in d.part.rels.values()}
    body_text = "\n".join(p.text for p in d.paragraphs)
    assert "[1]" in body_text and "[@S" not in body_text
    assert "PAGE" in d.sections[0].footer._element.xml  # noqa: SLF001
    assert d.core_properties.title == "Report a"


def test_docx_keeps_markup_as_text_and_tables(ac, g):
    section(
        g
    ).body = f"<b>x</b> and **bold** and a list:\n\n- one [@S{g['sources'][0].id}]\n- two\n\n| a | b |\n|---|---|\n| 1 | 2 |"
    db.session.commit()
    d = Document(io.BytesIO(get(ac, g, "docx").data))
    texts = [p.text for p in d.paragraphs]
    assert "<b>x</b> and bold and a list:" in texts
    assert any(p.style.name == "List Bullet" and p.text.startswith("one [1]") for p in d.paragraphs)
    assert [[c.text for c in row.cells] for row in d.tables[0].rows] == [["a", "b"], ["1", "2"]]


def test_pdf_export_has_pages_numbers_links_and_text(ac, g):
    r = get(ac, g, "pdf")
    assert r.mimetype == "application/pdf" and r.data.startswith(b"%PDF-")
    text, pdf = pdf_text(r.data)
    assert pdf.page_count >= 2 and pdf.metadata["title"] == "Report a"
    for needle in (
        "Report a",
        "Research scope",
        "Methodology",
        "Overview",
        "Quality assessment",
        "References",
        "Page 1",
        "Page 2",
    ):
        assert needle in text, needle
    flat = " ".join(text.split())
    assert "COP of 3 to 4 in mild weather [1]." in flat and "Heat pump guide a" in flat
    links = [lk for page in pdf for lk in page.get_links()]
    assert any(lk.get("uri") == "https://energy.gov/a0" for lk in links)  # clickable reference address
    assert any(lk["kind"] == fitz.LINK_GOTO for lk in links)  # citation -> reference


def test_pdf_has_contents_for_longer_reports(ac, alice):
    g2 = finish(build(alice, "long"), extra_sections=3)
    text, pdf = pdf_text(get(ac, g2, "pdf").data)
    assert "Contents" in text and "Extra 2" in text
    assert "Contents" in "\n".join(p.text for p in Document(io.BytesIO(get(ac, g2, "docx").data)).paragraphs)
    assert "## Contents" in get(ac, g2, "md").get_data(as_text=True)


def test_pdf_survives_markup_and_unsupported_scripts(ac, g):
    section(
        g
    ).body = '<b>unclosed </para> <a href="x">y</a> & more \u0627\u0631\u062f\u0648 \u4e2d\u6587 caf\u00e9'
    db.session.commit()
    r = get(ac, g, "pdf")
    assert r.status_code == 200
    text = " ".join(pdf_text(r.data)[0].split())
    assert "<b>unclosed </para>" in text and "caf\u00e9" in text
    assert "could not be drawn with the PDF font" in text


def test_exports_are_owner_only_and_logged(client, bob, ac, g):
    from sqlalchemy import select as sel

    assert get(ac, g, "pdf").status_code == 200
    assert get(ac, g, "rtf").status_code == 404
    logged = db.session.scalars(sel(AuditEvent).where(AuditEvent.event_type == "report_exported")).all()
    assert len(logged) == 1 and logged[0].details == {"format": "pdf"}
    ac.get("/auth/logout")
    other = client.application.test_client()
    login(other, "bob@example.com")
    for fmt in exports.FORMATS:
        assert other.get(url(g, f"/report/export/{fmt}")).status_code == 404
    assert other.get(url(g, "/report/quality")).status_code == 404
    anon = client.application.test_client()
    assert anon.get(url(g, "/report/export/pdf")).status_code in (302, 401)


def test_title_cannot_break_the_download_filename(ac, g):
    g["report"].title = 'a/../b"\r\nX-Evil: 1'
    db.session.commit()
    header = get(ac, g, "md").headers["Content-Disposition"]
    assert "\n" not in header and "\r" not in header and "/" not in header and header.count('"') == 2


def test_report_pages_offer_exports_and_quality(ac, g):
    html = ac.get(url(g, "/report")).get_data(as_text=True)
    for fmt in exports.FORMATS:
        assert f"/report/export/{fmt}" in html
    assert "/report/quality" in html
